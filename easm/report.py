"""Excel report: Summary pivot + Issues + per-type tabs (sheet-1 layout) + per-domain tabs (sheet-2 layout)."""
import datetime as dt
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import __version__
from .findings import SEV_ORDER

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
SEV_FILL = {"Critical": "C00000", "High": "FF6B6B", "Medium": "FFC000", "Low": "FFE699", "Info": "DDEBF7"}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
SEVS = ["Critical", "High", "Medium", "Low", "Info"]


def _table(ws, headers, rows, widths=None, row0=1, col0=1, sev_col=None):
    for j, h in enumerate(headers):
        c = ws.cell(row=row0, column=col0 + j, value=h)
        c.fill, c.font, c.border, c.alignment = HDR_FILL, HDR_FONT, BORDER, Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r):
            c = ws.cell(row=row0 + i, column=col0 + j, value=v)
            c.alignment, c.border = WRAP, BORDER
        if sev_col is not None:
            sc = ws.cell(row=row0 + i, column=col0 + sev_col)
            fill = SEV_FILL.get(str(sc.value))
            if fill:
                sc.fill = PatternFill("solid", fgColor=fill)
                if sc.value in ("Critical",):
                    sc.font = Font(color="FFFFFF", bold=True)
    for j, w in enumerate(widths or []):
        ws.column_dimensions[get_column_letter(col0 + j)].width = w
    ws.freeze_panes = ws.cell(row=row0 + 1, column=col0)
    if rows:
        ws.auto_filter.ref = f"{get_column_letter(col0)}{row0}:{get_column_letter(col0 + len(headers) - 1)}{row0 + len(rows)}"


def _safe_title(name, used):
    t = re.sub(r"[\[\]:*?/\\]", "", name)[:31] or "sheet"
    base, n = t, 2
    while t.lower() in used:
        t = f"{base[:28]}_{n}"; n += 1
    used.add(t.lower())
    return t


def _sorted(fs):
    return sorted(fs, key=lambda f: (f.domain, SEV_ORDER.get(f.severity, 9), f.host, f.name))


def build(ctx, fixed, path, period=None):
    wb = Workbook()
    used = set()
    d = ctx.data
    findings = _sorted(ctx.findings)
    period = period or dt.datetime.now().strftime("%B %Y")

    # ---- Cover
    ws = wb.active
    ws.title = _safe_title("Cover Page", used)
    ws.sheet_view.showGridLines = False
    ws["B3"] = "External Attack Surface Monitoring Report"; ws["B3"].font = Font(size=22, bold=True, color="1F3864")
    ws["B5"] = ctx.org; ws["B5"].font = Font(size=16, bold=True)
    ws["B6"] = f"Period: {period}"
    ws["B7"] = f"Run date: {ctx.started:%Y-%m-%d %H:%M}"
    ws["B8"] = f"Domains in scope: {len(ctx.roots)}"
    ws["B9"] = f"Tool: easm {__version__}"
    ws["B11"] = "Open findings by severity"; ws["B11"].font = Font(bold=True)
    for i, s in enumerate(SEVS):
        ws.cell(row=12 + i, column=2, value=s).fill = PatternFill("solid", fgColor=SEV_FILL[s])
        ws.cell(row=12 + i, column=3, value=sum(1 for f in findings if f.severity == s))
    ws["B18"] = "New this run"; ws["C18"] = sum(1 for f in findings if f.extra.get("status") == "New")
    ws["B19"] = "Fixed since last run"; ws["C19"] = len(fixed)
    ws["B21"] = "Confidential - prepared for authorized security assessment only."; ws["B21"].font = Font(italic=True, color="7F7F7F")
    ws.column_dimensions["B"].width = 34; ws.column_dimensions["C"].width = 14

    # ---- Summary (live COUNTIFS so edits to Issues update it, like a pivot)
    ws = wb.create_sheet(_safe_title("Summary", used))
    ws["A1"] = "Open findings by domain and severity (excludes Fixed)"; ws["A1"].font = Font(bold=True, size=12)
    hdr = ["Domains"] + SEVS + ["Grand Total"]
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(row=3, column=j, value=h); c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
    for i, dom in enumerate(ctx.roots, start=4):
        ws.cell(row=i, column=1, value=dom).border = BORDER
        for j, s in enumerate(SEVS, start=2):
            col = get_column_letter(j)
            ws.cell(row=i, column=j, value=f'=COUNTIFS(Issues!$B:$B,$A{i},Issues!$F:$F,{col}$3,Issues!$J:$J,"<>Fixed")').border = BORDER
        ws.cell(row=i, column=7, value=f"=SUM(B{i}:F{i})").border = BORDER
    last = 3 + len(ctx.roots)
    ws.cell(row=last + 1, column=1, value="Grand Total").font = Font(bold=True)
    for j in range(2, 8):
        col = get_column_letter(j)
        c = ws.cell(row=last + 1, column=j, value=f"=SUM({col}4:{col}{last})"); c.font = Font(bold=True); c.border = BORDER
    ws.column_dimensions["A"].width = 38
    for j in range(2, 8):
        ws.column_dimensions[get_column_letter(j)].width = 12

    # ---- Issues
    ws = wb.create_sheet(_safe_title("Issues", used))
    rows = []
    n = 0
    for f in findings:
        n += 1
        rows.append([n, f.domain, f.host, f.category, f.name, f.severity, f.remediation, "", "", f.extra.get("status", "Open")])
    for p in sorted(fixed, key=lambda p: (p["domain"], p["host"])):
        n += 1
        rows.append([n, p["domain"], p["host"], p["category"], p["name"], p["severity"], p["remediation"], "", "", "Fixed"])
    _table(ws, ["S:NO", "Domain", "Host", "Type", "Issues", "Severity", "Recommendation", "Correction", "Corrective action", "Status"],
           rows, [7, 26, 32, 16, 48, 11, 60, 18, 22, 10], sev_col=5)

    # ---- sheet-1 style tabs
    def by(*cats):
        return [f for f in findings if f.category in cats]

    ws = wb.create_sheet(_safe_title("DMARC & SPF", used))
    dns = d.get("dns_records", {})
    _table(ws, ["Domain", "Name", "Info", "Record", "Remediation", "Severity", "Remarks"],
           [[f.domain, f.name, f.description, f.evidence, f.remediation, f.severity, ""] for f in by("DMARC", "SPF", "DNS")],
           [24, 36, 50, 50, 50, 11, 24], sev_col=5)

    def vuln_rows(fs):
        return [[f.domain, f.url, f.host, f.name, f.description, f.remediation, f.reference, f.port, f.severity, f.evidence, ""] for f in fs]
    VH = ["Domain", "Url", "Subdomain", "Name", "Description", "Remediation", "Reference", "Port", "Severity", "Evidence", "Remarks"]
    VW = [22, 36, 30, 36, 50, 50, 36, 8, 11, 30, 22]
    for title, cats in (("SSL Vulnerability", ("SSL",)), ("Misconfiguration", ("Headers", "Misconfiguration")), ("Other Vulnerability", ("Other",))):
        ws = wb.create_sheet(_safe_title(title, used))
        _table(ws, VH, vuln_rows(by(*cats)), VW, sev_col=8)

    ws = wb.create_sheet(_safe_title("Open Ports", used))
    prow = []
    for ip, v in (d.get("ports", {}).get("ips", {}) or {}).items():
        for p in v["ports"]:
            prow.append([ip, ", ".join(v["hosts"]), p["port"], p["proto"], p["service"], f"{p['product']} {p['version']} {p['extrainfo']}".strip()])
    _table(ws, ["IP", "Hosts", "Port", "Proto", "Service", "Product / Version"], prow, [18, 44, 8, 8, 16, 44])

    ws = wb.create_sheet(_safe_title("Leaked Credentials", used))
    lk = d.get("leaks", {})
    _table(ws, ["Domain", "Leak Name", "Leak Description", "Leak Date", "Leak Data Classes", "Breached Email"],
           [[r["domain"], r["leak"], r["description"], r["date"], r["classes"], r["email"]] for r in lk.get("rows", [])] or
           [["NA"] * 6], [22, 24, 60, 12, 40, 34])
    if lk.get("skipped"):
        ws["H1"] = "Skipped: set HIBP_KEY (domain must be verified in HIBP) to enable."

    disc = d.get("discovery", {})
    ws = wb.create_sheet(_safe_title("Similar Domain", used))
    _table(ws, ["Domain", "Similar Domain", "Registrant Name", "Registrant Organization Name"],
           [[r["domain"], r["similar"], r["registrant_name"], r["registrant_org"]] for r in disc.get("similar", [])], [24, 32, 30, 40])
    ws = wb.create_sheet(_safe_title("Similar Domain SSL Cert", used))
    _table(ws, ["Domain", "Similar Domain", "Issuer", "Subject", "Not Before", "Not After"],
           [[r["domain"], r["similar"], r["cert"]["issuer"], r["cert"]["subject"], r["cert"]["not_before"][:19], r["cert"]["not_after"][:19]]
            for r in disc.get("similar", []) if r.get("cert")], [24, 30, 50, 36, 20, 20])
    ws = wb.create_sheet(_safe_title("Discovered Domains", used))
    _table(ws, ["Candidate domain (NOT in scope - verify ownership)", "Matched keyword(s)"],
           [[k, ", ".join(v)] for k, v in sorted(disc.get("candidates", {}).items())], [52, 28])

    ws = wb.create_sheet(_safe_title("Emails", used))
    em = d.get("subdomains", {}).get("emails", {})
    _table(ws, ["Domain", "Email", "Status"], [[dom, e, "Found (public)"] for e, dom in sorted(em.items())] or [["NA", "NA", "NA"]], [26, 44, 18])

    live = d.get("liveness", [])
    best = {}
    for r in live:
        if r["host"] not in best or r["proto"] == "https":
            best[r["host"]] = r
    hosts = d.get("subdomains", {}).get("hosts", {})
    ws = wb.create_sheet(_safe_title("Subdomains", used))
    _table(ws, ["Domain", "Host", "Resolves", "IPs", "CNAME", "Live", "HTTP code", "Title", "Server"],
           [[v["root"], h, "Yes" if v["resolves"] else "No", ", ".join(v["ips"]), ", ".join(v["cname"]),
             "Yes" if h in best else "No", best[h]["code"] if h in best else "", best[h]["title"] if h in best else "",
             best[h]["server"] if h in best else ""] for h, v in sorted(hosts.items(), key=lambda kv: (kv[1]["root"], kv[0]))],
           [24, 40, 9, 30, 36, 6, 10, 40, 22])

    # ---- sheet-2 style: one tab per domain
    tls = d.get("webchecks", {}).get("tls", {})
    for dom in ctx.roots:
        ws = wb.create_sheet(_safe_title(dom, used))
        ws["A1"] = dom; ws["A1"].font = Font(bold=True, size=13)
        dom_live = sorted(h for h in best if hosts.get(h, {}).get("root") == dom)
        ws["A3"] = "Web sites present." if dom_live else "No Website."
        r = 4
        if dom_live:
            rows = []
            for h in dom_live:
                hf = [f for f in findings if f.host == h and f.category == "Headers" and f.severity != "Info"]
                sf = [f for f in findings if f.host == h and f.category == "SSL"]
                resp = best[h]
                hdr = "NA" if resp["code"] >= 500 or resp["code"] == 0 else ("Fail" if hf else "Pass")
                ssl_ = "NA" if h not in tls or "error" in tls.get(h, {}) else ("Fail" if sf else "Pass")
                rows.append([h, "Yes", hdr, ssl_])
            _table(ws, ["Active sub-domains:", "Public access", "Security Headers", "SSL Scan"], rows, [40, 16, 18, 14], row0=r)
            for rr in range(r + 1, r + 1 + len(rows)):
                for cc in (3, 4):
                    c = ws.cell(row=rr, column=cc)
                    c.fill = PatternFill("solid", fgColor={"Pass": "C6EFCE", "Fail": "FFC7CE"}.get(c.value, "EDEDED"))
            r += len(rows) + 2
        ws.cell(row=r, column=1, value="Records:").font = Font(bold=True); r += 1
        rec = [f for f in findings if f.domain == dom and f.category in ("DMARC", "SPF", "DNS")]
        if rec:
            for f in rec:
                ws.cell(row=r, column=1, value=f"{f.name} [{f.severity}]"); r += 1
        else:
            ws.cell(row=r, column=1, value="All are good."); r += 1
        r += 1
        ws.cell(row=r, column=1, value="Comments:").font = Font(bold=True); r += 1
        rr = dns.get(dom, {})
        notes = []
        if rr.get("mx"): notes.append(f"MX: {', '.join(rr['mx'])}")
        if not rr.get("mx"): notes.append("No MX records (domain does not receive mail).")
        if rr.get("dkim_selectors"): notes.append("DKIM selectors found: " + ", ".join(rr["dkim_selectors"]))
        risky = [f for f in findings if f.domain == dom and f.category == "Ports"]
        for f in risky: notes.append(f"Exposed: {f.name} on {f.host}")
        for n_ in notes:
            ws.cell(row=r, column=1, value=n_); r += 1
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 60)

    wb.save(path)
    return path

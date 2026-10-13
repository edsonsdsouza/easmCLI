"""Offline end-to-end test: mocks DNS/HTTP/TLS so the whole pipeline + xlsx runs without network."""
import datetime as dt
import os
import sys
import tempfile

from openpyxl import load_workbook

from easm import resolver
from easm.stages import dns_records, discovery, leaks, liveness, ports, subdomains, webchecks
from easm.cli import main

ZONE = {
    ("acme-test.com", "NS"): ["ns1.acme-test.com."], ("acme-test.com", "MX"): ["10 mail.acme-test.com."],
    ("acme-test.com", "A"): ["93.184.216.34"],
    ("acme-test.com", "TXT"): ['"v=spf1 include:_spf.google.com ~all"'],
    ("_dmarc.acme-test.com", "TXT"): ['"v=DMARC1; p=none; rua=mailto:x@acme-test.com"'],
    ("www.acme-test.com", "A"): ["93.184.216.34"], ("app.acme-test.com", "A"): ["93.184.216.35"],
    ("old.acme-test.com", "CNAME"): ["gone.example-cloud.net."],
    ("other-test.org", "NS"): ["ns1.other-test.org."], ("other-test.org", "A"): [],
}
NX = {"gone.example-cloud.net"}


def fake_query(res, name, rtype):
    name = name.rstrip(".")
    if (name, rtype) in ZONE:
        return ZONE[(name, rtype)]
    if rtype == "TXT" and name.endswith("acme-test.com") or name.endswith("other-test.org"):
        return []
    return []


def fake_nx(res, name):
    return name.rstrip(".") in NX


def fake_probe(host, proto):
    base = {"host": host, "proto": proto, "url": f"{proto}://{host}"}
    if host == "www.acme-test.com":
        if proto == "https":
            return {**base, "code": 200, "final_url": f"https://{host}/", "title": "Acme", "server": "nginx/1.18.0",
                    "headers": {"server": "nginx/1.18.0", "strict-transport-security": "max-age=1"},
                    "set_cookie": ["sid=1; Path=/"], "redirected_https": True}
        return {**base, "code": 200, "final_url": f"http://{host}/", "title": "Acme", "server": "nginx",
                "headers": {}, "set_cookie": [], "redirected_https": False}
    if host == "app.acme-test.com" and proto == "https":
        return {**base, "code": 200, "final_url": f"https://{host}/", "title": "App", "server": "",
                "headers": {"content-security-policy": "default-src 'self'", "strict-transport-security": "max-age=31536000",
                            "x-content-type-options": "nosniff", "referrer-policy": "no-referrer"}, "set_cookie": [], "redirected_https": True}
    return {**base, "code": 0}


def fake_tls(host, port=443):
    exp = (dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=5)).isoformat() if host.startswith("www") else \
          (dt.datetime.now(dt.timezone.utc) + dt.timedelta(days=200)).isoformat()
    return {"host": host, "port": port, "not_after": exp, "not_before": exp, "issuer": "CN=Test CA", "subject": f"CN={host}",
            "names": [host], "old_protocols": ["TLSv1.0"] if host.startswith("app") else []}


def run():
    resolver.query = fake_query
    resolver.nxdomain = fake_nx
    for mod in (dns_records, subdomains, discovery):
        mod.query = fake_query
    subdomains.nxdomain = fake_nx
    subdomains.crtsh = lambda d: {"www.acme-test.com", "app.acme-test.com", "old.acme-test.com"} if d == "acme-test.com" else set()
    subdomains.hackertarget = subdomains.otx = lambda d: set()
    liveness.probe = fake_probe
    webchecks.tls_check = fake_tls
    dns_records._axfr = lambda ns, d: None
    with tempfile.TemporaryDirectory() as tmp:
        argv = ["run", "-d", "acme-test.com", "-d", "other-test.org", "--org", "Acme", "--skip", "discovery", "--skip", "ports",
                "--out-root", tmp]
        try:
            main(argv)
        except SystemExit as e:
            assert e.code == 0, e.code
        run_dir = [os.path.join(tmp, d) for d in os.listdir(tmp)][0]
        # second run -> diff against the first
        import time; time.sleep(61)  # distinct timestamp folder
        webchecks.tls_check = lambda h, port=443: {**fake_tls(h), "not_after": (dt.datetime.now(dt.timezone.utc) + dt.timedelta(days=300)).isoformat()}
        try:
            main(argv)
        except SystemExit as e:
            assert e.code == 0
        runs = sorted(os.path.join(tmp, d) for d in os.listdir(tmp))
        assert len(runs) == 2, runs
        import json
        f1 = json.load(open(os.path.join(runs[0], "findings.json")))
        f2 = json.load(open(os.path.join(runs[1], "findings.json")))
        checks1 = {f["check"] for f in f1}
        print("run1 checks:", sorted(checks1))
        for want in ("dmarc_none", "dmarc_sp", "tls_expired", "tls_old_proto", "hdr_csp", "hdr_xfo", "cookie_httponly", "cookie_secure", "http_no_redirect", "dangling_cname", "hdr_server", "dmarc_missing"):
            assert want in checks1, f"missing {want}"
        assert not any(f["check"] == "spf_missing" and f["domain"] == "acme-test.com" for f in f1)
        assert any(f["check"] == "spf_missing" and f["domain"] == "other-test.org" for f in f1)
        assert "tls_expired" not in {f["check"] for f in f2}
        fixed = json.load(open(os.path.join(runs[1], "fixed.json")))
        assert any(x["check"] == "tls_expired" for x in fixed), "expired cert should show as Fixed"
        assert all(f["extra"]["status"] == "Open" for f in f2)
        xl = [os.path.join(runs[1], x) for x in os.listdir(runs[1]) if x.endswith(".xlsx")][0]
        wb = load_workbook(xl)
        print("sheets:", wb.sheetnames)
        iss = wb["Issues"]
        statuses = {r[9] for r in iss.iter_rows(min_row=2, values_only=True)}
        assert "Fixed" in statuses and "Open" in statuses, statuses
        print("Issues rows:", iss.max_row - 1, "statuses:", statuses)
        dest = os.environ.get("EASM_TEST_COPY")
        if dest:
            import shutil; shutil.copy(xl, dest)
    print("OK")


if __name__ == "__main__":
    run()


def test_public_dns():
    """Name lookups must come from public resolvers only; internal-only names must NOT resolve."""
    import socket
    from easm import resolver, util
    resolver.query = lambda res, name, rtype: ["203.0.113.7"] if (name == "public.test" and rtype == "A") else ([] if rtype in ("A", "AAAA") else None)
    util.use_public_dns(["8.8.8.8", "1.1.1.1"])
    got = socket.getaddrinfo("public.test", 443, socket.AF_INET, socket.SOCK_STREAM)
    assert got[0][4][0] == "203.0.113.7", got
    try:
        socket.getaddrinfo("internal-only.corp", 443)
        raise AssertionError("internal name resolved")
    except socket.gaierror:
        pass
    assert socket.getaddrinfo("127.0.0.1", 80)  # IP literals untouched
    socket.getaddrinfo = util._orig_getaddrinfo
    print("public DNS OK")
